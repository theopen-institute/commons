"""A bank statement, read into rows that can become Bank Transactions.

`start_reading` takes the file and hands it to a background job, and
`reading_status` reports on that job until it answers with the rows it found
and what the statement says about itself: the account number, the opening and
closing balances, notes worth a look. Nothing is written. Which
rows become Bank Transactions is decided in the browser, where the rows are
checked against the account's existing lines and imported through
`frappe.client.insert_many`, which inserts and submits each one through its
own permission check. See `frontend/src/data/statementImport.ts`.

Two ways of reading, chosen by what the file is
-----------------------------------------------
**A spreadsheet** (XLSX, XLS or CSV) is already rows and columns. The only hard
part is knowing which column is which, and banks export them every which way:
one amount column with a sign, one with Dr/Cr beside it, separate debit and
credit columns, a header on row 7 below the bank's address. So Claude is shown
the first rows (`SAMPLE_ROWS`) and asked only for the layout (`MAPPING`), and
the rows are then read by `apply_mapping`, here, in code. A statement of a
thousand lines costs one small call, and no figure is ever copied by a model.

**A PDF or an image** has no columns to map, so Claude copies the rows
themselves (`ROWS`), through `commons.api_integrations.claude.documents.read`,
the same call Document Capture reads invoices with. Copying is the one thing
a model can get subtly wrong, which is why the rows carry the running balance
where the statement prints one: the browser walks the balances and points at
any row where they stop adding up.

Dates in both cases come back as printed, with year, month, day and calendar,
because Nepali statements are sometimes dated in Bikram Sambat. The browser
converts them with the tables Document Capture uses.

Why a background job
--------------------
Copying a long PDF row by row can take Claude several minutes, and a web
request is cut off at two minutes on production (gunicorn's timeout). So the
reading runs in the `long` queue, where it may take `JOB_TIMEOUT`, and Claude's
answer is streamed so the job can count the rows as they arrive. The job's
state lives in the cache under a token only the person who started it may
read: queued, reading (with rows so far), done (with the result) or failed
(with a sentence). The file waits in the cache too, and is dropped as soon as
the job has read it. That plumbing, and how a job RQ killed is told from one
still running, is `commons.api_integrations.claude.jobs`, which Document
Capture's invoice reading shares.

Who may use it
--------------
Whoever may create a Bank Transaction, checked before anything is sent to
Claude, because every read is billed. There is also a per-person limit of
`HOURLY_LIMIT` reads an hour, as Document Capture has.
"""

import csv
import datetime
import io
import re

import frappe
from frappe import _

from commons.api_integrations.claude import client as claude
from commons.api_integrations.claude import documents, jobs
from commons.commons_core import apps

BANK_TRANSACTION = "Bank Transaction"

HOURLY_LIMIT = 30

# How much of a spreadsheet Claude sees to work out its layout. Enough to get
# past a bank's letterhead and a header row and into the transactions.
SAMPLE_ROWS = 40
SAMPLE_CELL = 80

# Seconds. How long Claude may take over one statement in the background,
# and a margin over it for the job as a whole.
READ_TIMEOUT = 600
JOB_TIMEOUT = READ_TIMEOUT + 120

# How long a finished or abandoned reading is kept for the dialog to collect.
STATE_TTL = jobs.STATE_TTL

# A copied statement's answer can be long: a hundred rows is ten thousand
# tokens or so. Streaming allows far more than the 16,000 of a web request.
MAX_TOKENS = 64000

XLSX = "xlsx"
XLS = "xls"
CSV = "csv"
DOCUMENT = "document"

READING = jobs.Job("statement_import", "commons.banking.statement_import.run_reading", JOB_TIMEOUT, STATE_TTL)


def available() -> bool:
	"""Whether statements can be imported here: bank transactions to import
	into, and a Claude key to read with."""
	return apps.has_doctype(BANK_TRANSACTION) and claude.available()


def can_import() -> bool:
	return bool(frappe.has_permission(BANK_TRANSACTION, "create"))


@frappe.whitelist()
def import_available() -> bool:
	"""Whether the page should offer the import at all, for this reader."""
	return available() and can_import()


def _require() -> None:
	if not available():
		frappe.throw(_("Importing bank statements is not set up on this site."))
	if not can_import():
		frappe.throw(_("You are not allowed to create bank transactions."), frappe.PermissionError)


# --------------------------------------------------------------------------- #
# Schemas                                                                      #
# --------------------------------------------------------------------------- #


def _nullable(schema: dict, description: str | None = None) -> dict:
	nullable = {"anyOf": [schema, {"type": "null"}]}
	if description:
		nullable["description"] = description
	return nullable


def _object(properties: dict) -> dict:
	return {
		"type": "object",
		"properties": properties,
		"required": list(properties),
		"additionalProperties": False,
	}


STRING = {"type": "string"}
NUMBER = {"type": "number"}
INTEGER = {"type": "integer"}
CALENDAR = {
	"type": "string",
	"enum": ["AD", "BS"],
	"description": "BS for Bikram Sambat dates, AD for Gregorian ones.",
}

DATE = _object(
	{
		"printed": {"type": "string", "description": "The date exactly as printed."},
		"year": INTEGER,
		"month": {"type": "integer", "description": "1 to 12."},
		"day": INTEGER,
		"calendar": CALENDAR,
	}
)

# What a statement says about itself, whichever way it was read.
HEADER = {
	"is_statement": {"type": "boolean", "description": "Whether this is a bank statement at all."},
	"account_number": _nullable(STRING, "The account number the statement is for, as printed."),
	"currency": _nullable(STRING, "ISO 4217 code."),
	"opening_balance": _nullable(NUMBER, "The balance brought forward, as printed."),
	"closing_balance": _nullable(NUMBER, "The balance at the end, as printed."),
	"notes": {"type": "array", "items": STRING},
}

ROWS = _object(
	{
		**HEADER,
		"rows": {
			"type": "array",
			"items": _object(
				{
					"date": DATE,
					"description": STRING,
					"reference": _nullable(STRING, "Cheque number or transaction reference, if in its own column."),
					"withdrawal": {"type": "number", "description": "Money out of the account, or 0."},
					"deposit": {"type": "number", "description": "Money into the account, or 0."},
					"balance": _nullable(NUMBER, "The running balance printed on this row, if any."),
				}
			),
		},
	}
)

COLUMN = _nullable(INTEGER)

MAPPING = _object(
	{
		**HEADER,
		"first_data_row": {"type": "integer", "description": "Index of the first transaction row."},
		"date_column": INTEGER,
		"date_order": {
			"type": "string",
			"enum": ["DMY", "MDY", "YMD"],
			"description": "The order of day, month and year in dates written as text.",
		},
		"calendar": CALENDAR,
		"description_columns": {"type": "array", "items": INTEGER},
		"reference_column": COLUMN,
		"withdrawal_column": _nullable(INTEGER, "A column holding only money out."),
		"deposit_column": _nullable(INTEGER, "A column holding only money in."),
		"amount_column": _nullable(INTEGER, "A single column holding both, when there is no pair."),
		"amount_sign": _nullable(
			{"type": "string", "enum": ["negative_is_withdrawal", "positive_is_withdrawal"]},
			"How the single amount column says which way the money went, when there is no Dr/Cr column.",
		),
		"direction_column": _nullable(INTEGER, "A column saying Dr/Cr or Debit/Credit beside a single amount."),
		"balance_column": COLUMN,
	}
)

COMMON_RULES = """\
Copy, don't compute: every figure must be one printed in the statement. Numbers \
are plain numbers, without currency symbols or thousands separators; convert \
Devanagari digits (०१२३४५६७८९) to 0-9; South Asian grouping such as 1,13,000.00 \
is 113000.

Direction is from the account holder's point of view. A withdrawal is money out \
of the account (a debit on the statement: payments, transfers out, fees, \
cheques paid); a deposit is money in (a credit: receipts, transfers in, \
interest). Banks print debits and credits from their own ledger, so "Dr" means \
money out and "Cr" means money in.

Dates: Nepali statements are sometimes dated in Bikram Sambat (B.S. or वि.सं.), \
whose years currently run from about 2075 to 2090. Mark those "BS" and do not \
convert them. Gregorian dates are "AD".

currency: NPR for Nepali rupees, INR, USD and so on; null if it cannot be told. \
If the document is not a bank statement, set is_statement to false. notes: \
anything a bookkeeper should check, such as pages that seem to be missing, \
figures you are unsure of, or balances that do not add up; a sentence each, \
and an empty list if there is nothing to say."""

DOCUMENT_INSTRUCTIONS = f"""\
This is a bank statement for one of our organisation's bank accounts. Copy every \
transaction row into the schema, in the order printed.

{COMMON_RULES}

- Leave out rows that are not transactions: balance brought forward, page \
totals, closing balance, headers repeated on each page. Put the brought-forward \
and closing balances in opening_balance and closing_balance instead.
- description is the row's narration as printed, all of it. reference is a \
cheque or reference number only where the statement prints it in its own \
column; otherwise null.
- balance is the running balance printed on the row, if the statement has that \
column, else null."""

MAPPING_INSTRUCTIONS = f"""\
These are the first rows of a bank statement exported as a spreadsheet, for one \
of our organisation's bank accounts. Each line is one row: its index, then its \
cells separated by " | ", each cell preceded by its column index. Work out the \
layout so the rest of the rows can be read by a program.

{COMMON_RULES}

- first_data_row is the index of the first transaction, after any letterhead, \
header row, or balance brought forward.
- Give the amounts either as a withdrawal_column and a deposit_column, or as one \
amount_column with either a direction_column (Dr/Cr) or an amount_sign. Leave \
the others null.
- description_columns are the columns whose text together makes the \
narration, in order.
- opening_balance, closing_balance and account_number only if they appear in \
these rows."""


# --------------------------------------------------------------------------- #
# The endpoint                                                                 #
# --------------------------------------------------------------------------- #


@frappe.whitelist(methods=["POST"])
def start_reading() -> dict:
	"""Take the statement in the request's `file` field and start reading it.

	A multipart upload, as Document Capture's `start_reading` takes. Answers at
	once with a token for `reading_status`.
	"""
	_require()
	upload = frappe.request.files.get("file") if frappe.request else None
	if not upload:
		frappe.throw(_("Choose a statement to read."))
	content = upload.stream.read()
	if len(content) > documents.MAX_BYTES:
		frappe.throw(
			_("That file is over {0} MB. Send a smaller one.").format(documents.MAX_BYTES // (1024 * 1024))
		)
	file_kind(content)
	_count_read()
	return {"token": READING.start(content, rows=0, step=None)}


@frappe.whitelist()
def reading_status(token: str) -> dict:
	"""Where a reading has got to. Only its starter may ask."""
	return READING.status(token, lost=_lost())


def run_reading(token: str) -> None:
	"""The background job. Runs as the user who started it, since
	`frappe.enqueue` carries the session user into the job."""
	READING.run(
		token,
		lambda content, progress: read_statement(content, progress=progress),
		missing=_("The statement was not found. Start again."),
		failed=_("Something went wrong reading the statement."),
		lost=_lost(),
		log_title="Bank statement import failed",
	)


def _lost() -> str:
	return _(
		"The reading stopped without an answer. It may have run out of time, or the worker stopped. Try again."
	)


def read_statement(content: bytes, progress=lambda **changes: None) -> dict:
	"""Everything the dialog needs from one statement file.

	`progress` is told what is happening as it happens: the step, and while
	Claude copies a PDF, how many rows it has written so far.
	"""
	kind = file_kind(content)
	if kind == DOCUMENT:
		progress(status="reading", step="copying")
		counter = RowCounter(lambda rows: progress(rows=rows))
		read = documents.read(
			content,
			ROWS,
			DOCUMENT_INSTRUCTIONS,
			on_text=counter.feed,
			timeout=READ_TIMEOUT,
			max_tokens=MAX_TOKENS,
			too_long=_(
				"That statement is too long to read in one go. Send fewer pages at a time, or the bank's spreadsheet export."
			),
		)
		rows = [_document_row(row) for row in read.get("rows") or []]
	else:
		progress(status="reading", step="columns")
		grid = read_grid(content, kind)
		if not grid:
			frappe.throw(_("That spreadsheet has no rows in it."))
		read = _read_mapping(grid)
		progress(step="rows")
		read["notes"] = list(read.get("notes") or [])
		rows = apply_mapping(grid, read, notes=read["notes"])
		progress(rows=len(rows))

	return {
		"source": "spreadsheet" if kind != DOCUMENT else "document",
		"is_statement": bool(read.get("is_statement", True)),
		"account_number": read.get("account_number"),
		"currency": read.get("currency"),
		"opening_balance": read.get("opening_balance"),
		"closing_balance": read.get("closing_balance"),
		"notes": read.get("notes") or [],
		"rows": rows,
		"model": claude.model(),
	}


# Counts a copied statement's rows as they stream in: each has a description.
RowCounter = jobs.RowCounter


def _state_key(token: str) -> str:
	return READING.state_key(token)


def _file_key(token: str) -> str:
	return READING.file_key(token)


def _count_read() -> None:
	key = frappe.cache.make_key(f"commons:statement_import:reads:{frappe.session.user}")
	count = frappe.cache.incrby(key, 1)
	if count == 1:
		frappe.cache.expire(key, 3600)
	if count > HOURLY_LIMIT:
		frappe.throw(
			_("That is {0} statements in the last hour, which is the limit. Try again later.").format(
				HOURLY_LIMIT
			),
			frappe.RateLimitExceededError,
		)


def _document_row(row: dict) -> dict:
	return {
		"date": row.get("date"),
		"description": (row.get("description") or "").strip(),
		"reference": (row.get("reference") or "").strip() or None,
		"withdrawal": abs(float(row.get("withdrawal") or 0)),
		"deposit": abs(float(row.get("deposit") or 0)),
		"balance": row.get("balance"),
	}


# --------------------------------------------------------------------------- #
# Spreadsheets                                                                 #
# --------------------------------------------------------------------------- #


def file_kind(content: bytes) -> str:
	"""What the file is, by its bytes. The name and the browser's type are
	only claims."""
	if not content:
		frappe.throw(_("That file is empty."))
	if content.startswith(b"%PDF"):
		return DOCUMENT
	if content.startswith(b"PK\x03\x04"):
		return XLSX
	if content.startswith(b"\xd0\xcf\x11\xe0"):
		return XLS
	if _text(content) is not None:
		return CSV
	return DOCUMENT


def _text(content: bytes) -> str | None:
	"""The file as text, if it is text with the look of a table."""
	for encoding in ("utf-8-sig", "cp1252"):
		try:
			text = content.decode(encoding)
		except UnicodeDecodeError:
			continue
		sample = text[:4000]
		if "\x00" in sample:
			return None
		if any(delimiter in sample for delimiter in (",", ";", "\t")) and "\n" in sample:
			return text
		return None
	return None


def read_grid(content: bytes, kind: str) -> list[list]:
	"""Every row of the first sheet as a list of cell values, trailing empty
	rows and columns trimmed. Dates from Excel arrive as `datetime`."""
	if kind == XLSX:
		from openpyxl import load_workbook

		sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
		rows = [list(row) for row in sheet.iter_rows(values_only=True)]
	elif kind == XLS:
		import xlrd

		book = xlrd.open_workbook(file_contents=content)
		sheet = book.sheet_by_index(0)
		rows = []
		for index in range(sheet.nrows):
			row = []
			for cell in sheet.row(index):
				if cell.ctype == xlrd.XL_CELL_DATE:
					row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
				else:
					row.append(cell.value)
			rows.append(row)
	else:
		text = _text(content) or ""
		try:
			dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
		except csv.Error:
			dialect = csv.excel
		rows = list(csv.reader(io.StringIO(text), dialect))

	rows = [[None if cell in ("", None) else cell for cell in row] for row in rows]
	while rows and not any(cell is not None for cell in rows[-1]):
		rows.pop()
	return rows


def sample(grid: list[list]) -> str:
	"""The first rows, the way `MAPPING_INSTRUCTIONS` describes them."""
	lines = []
	for index, row in enumerate(grid[:SAMPLE_ROWS]):
		cells = [
			f"{column}: {_cell_text(cell)[:SAMPLE_CELL]}"
			for column, cell in enumerate(row)
			if cell is not None
		]
		lines.append(f"{index}. " + " | ".join(cells))
	return "\n".join(lines)


def _cell_text(cell) -> str:
	if isinstance(cell, datetime.datetime):
		return cell.date().isoformat() if cell.time() == datetime.time() else cell.isoformat()
	if isinstance(cell, datetime.date):
		return cell.isoformat()
	if isinstance(cell, float) and cell.is_integer():
		return str(int(cell))
	return str(cell).strip()


def _read_mapping(grid: list[list]) -> dict:
	response = claude.create_message(
		max_tokens=4000,
		output_config={"effort": "medium", "format": {"type": "json_schema", "schema": MAPPING}},
		messages=[
			{"role": "user", "content": f"{MAPPING_INSTRUCTIONS}\n\n<rows>\n{sample(grid)}\n</rows>"}
		],
	)
	if response.stop_reason == "refusal":
		claude.fail(_("Claude declined to read this statement."))
	text = next((block.text for block in response.content if block.type == "text"), None)
	if not text:
		claude.fail(_("Claude sent back nothing to read."))
	return frappe.parse_json(text)


def apply_mapping(grid: list[list], mapping: dict, notes: list | None = None) -> list[dict]:
	"""Every transaction row of the sheet, read with the layout Claude gave.

	A row whose date cell is not a date is not a transaction (a page total, a
	closing balance, a footer) and is skipped, as is a row with no amount.

	A dated row whose amount is there but can't be read, or whose direction
	column says neither way, is left out too -- but said so in `notes`, by row.
	A missing row is something a bookkeeper can look for; a row read at a
	hundredth of its value, or booked the wrong way round, is not.
	"""
	rows = []
	unreadable, undirected = [], []

	def cell(row, column):
		return row[column] if column is not None and 0 <= column < len(row) else None

	first = max(0, int(mapping.get("first_data_row") or 0))
	d_is_deposit = _d_means_deposit(grid[first:], mapping.get("direction_column"), cell)
	for index, row in enumerate(grid[first:], start=first):
		date = parse_date(cell(row, mapping.get("date_column")), mapping.get("date_order") or "DMY")
		if not date:
			continue
		try:
			withdrawal, deposit = _amounts(row, mapping, cell, d_is_deposit)
		except UnreadableAmount:
			unreadable.append(index + 1)
			continue
		except UnknownDirection:
			undirected.append(index + 1)
			continue
		if not withdrawal and not deposit:
			continue
		description = " ".join(
			_cell_text(value)
			for value in (cell(row, column) for column in mapping.get("description_columns") or [])
			if value is not None
		)
		reference = cell(row, mapping.get("reference_column"))
		rows.append(
			{
				"date": {**date, "calendar": mapping.get("calendar") or "AD"},
				"description": description.strip(),
				"reference": _cell_text(reference) if reference is not None else None,
				"withdrawal": withdrawal,
				"deposit": deposit,
				"balance": parse_amount(cell(row, mapping.get("balance_column"))),
			}
		)

	if notes is not None:
		if unreadable:
			notes.append(
				_("Left out {0} dated rows whose amount could not be read (sheet rows {1}).").format(
					len(unreadable), _row_list(unreadable)
				)
			)
		if undirected:
			notes.append(
				_(
					"Left out {0} rows whose direction column says neither money in nor money out (sheet rows {1})."
				).format(len(undirected), _row_list(undirected))
			)
	return rows


class UnreadableAmount(Exception):
	"""A cell holds figures, but not in any shape an amount is printed in."""


class UnknownDirection(Exception):
	"""A direction cell says something other than money in or money out."""


def _row_list(numbers: list[int], shown: int = 10) -> str:
	listed = ", ".join(str(number) for number in numbers[:shown])
	return listed + (", …" if len(numbers) > shown else "")


def _figure(value) -> float | None:
	"""An amount cell's figure: None when it holds no digits at all -- blank, a
	dash, "nil" -- and `UnreadableAmount` when it holds some that aren't one."""
	if value is None or isinstance(value, bool):
		return None
	if isinstance(value, int | float):
		return float(value)
	if not re.search(r"\d", str(value).translate(DEVANAGARI)):
		return None
	amount = parse_amount(value)
	if amount is None:
		raise UnreadableAmount(value)
	return amount


# What a direction column says, reduced to its letters. A bare "D" and "C" are
# settled separately: see `_d_means_deposit`.
MONEY_OUT = {"dr", "db", "debit", "debited", "withdrawal", "withdrawals", "withdraw", "wdl", "wd", "w", "out"}
MONEY_IN = {"cr", "credit", "credited", "deposit", "deposits", "dep", "in"}
W_WORDS = {"w", "wd", "wdl", "withdrawal", "withdrawals", "withdraw"}


def _direction_word(value) -> str:
	return re.sub(r"[^a-z]", "", str(value).lower())


def _d_means_deposit(rows, column, cell) -> bool:
	"""Whether a bare "D" is a deposit rather than a debit.

	It is when the same column marks withdrawals with a W -- D/W is deposit and
	withdrawal -- and a debit otherwise, beside a C for credit.
	"""
	if column is None:
		return False
	return any(_direction_word(cell(row, column)) in W_WORDS for row in rows if cell(row, column) is not None)


def _amounts(row, mapping, cell, d_is_deposit: bool = False) -> tuple[float, float]:
	if mapping.get("withdrawal_column") is not None or mapping.get("deposit_column") is not None:
		out = _figure(cell(row, mapping.get("withdrawal_column"))) or 0
		into = _figure(cell(row, mapping.get("deposit_column"))) or 0
		# A bank that prints withdrawals as negatives in their own column.
		return abs(out), abs(into)

	amount = _figure(cell(row, mapping.get("amount_column")))
	if not amount:
		return 0, 0
	direction = cell(row, mapping.get("direction_column"))
	if direction is not None and str(direction).strip():
		word = _direction_word(direction)
		if word == "d":
			word = "deposit" if d_is_deposit else "dr"
		elif word == "c":
			word = "cr"
		if word in MONEY_OUT:
			return abs(amount), 0
		if word in MONEY_IN:
			return 0, abs(amount)
		# Guessing here books every such row one way or the other, silently.
		raise UnknownDirection(direction)
	raw = cell(row, mapping.get("amount_column"))
	marker = _direction_word(raw) if isinstance(raw, str) else ""
	if marker.endswith("dr"):
		return abs(amount), 0
	if marker.endswith("cr"):
		return 0, abs(amount)
	if mapping.get("amount_sign") == "positive_is_withdrawal":
		return (amount, 0) if amount > 0 else (0, -amount)
	return (-amount, 0) if amount < 0 else (0, amount)


DEVANAGARI = str.maketrans("०१२३४५६७८९", "0123456789")

# Minus signs that are not the ASCII hyphen: the Unicode minus, dashes, and the
# small and full-width hyphens.
MINUS = str.maketrans(dict.fromkeys("−‒–﹣－", "-"))

# What may stand beside a figure: a currency ("Rs.", "NPR", "₹", "रु.") or a Dr/Cr,
# at most this many characters once spaces, dots and a minus are taken out. Any
# longer and it is text with a number in it -- "Balance 5", "Page 2 of 3".
AFFIX = 4


def parse_amount(value) -> float | None:
	"""A figure as a statement prints it: `1,13,000.00`, `(500.00)`,
	`2,500.00 Dr`, `-75`, `१२००`, `Rs. 1,000.00`, `−500`, `1.234,56`. None for a
	blank, for text, and for anything that can't be read without guessing."""
	if value is None:
		return None
	if isinstance(value, bool):
		return None
	if isinstance(value, int | float):
		return float(value)
	text = str(value).translate(DEVANAGARI).translate(MINUS).strip()
	negative = False
	if text.startswith("(") and text.endswith(")"):
		negative, text = True, text[1:-1].strip()

	digits = [index for index, char in enumerate(text) if char.isdigit()]
	if not digits:
		return None
	prefix, figure, suffix = text[: digits[0]], text[digits[0] : digits[-1] + 1], text[digits[-1] + 1 :]
	for affix in (prefix, suffix):
		if "-" in affix:
			negative = True
		if len(re.sub(r"[\s.\-]", "", affix)) > AFFIX:
			return None

	number = _unseparated(figure)
	if number is None:
		return None
	return -number if negative else number


def _unseparated(figure: str) -> float | None:
	"""`figure` -- digits, and whatever separates them -- as a number.

	The decimal separator is the last of a comma and a dot when both appear
	(1,234.56 and 1.234,56). A lone comma is one only before one or two digits
	at the end (12,50): a comma before three is grouping, and South Asian
	grouping (1,13,000) never ends in two. Grouping must look like grouping --
	one to three digits, then groups of two or three, ending in three -- so a
	figure with two decimal points is refused rather than read as something.
	"""
	figure = re.sub(r"[\s' ]", "", figure)
	if not re.fullmatch(r"\d[\d,.]*", figure):
		return None
	commas, dots = figure.count(","), figure.count(".")
	if commas and dots:
		decimal = "," if figure.rfind(",") > figure.rfind(".") else "."
	elif commas:
		decimal = "," if commas == 1 and re.search(r",\d{1,2}$", figure) else None
	elif dots:
		decimal = "." if dots == 1 else None
	else:
		decimal = None

	whole, _sep, fraction = figure.rpartition(decimal) if decimal else (figure, "", "")
	if decimal and (not whole or decimal in whole):
		return None
	grouping = {",", "."} - {decimal}
	groups = re.split(r"[,.]", whole)
	if len(groups) > 1:
		if any(sep not in grouping for sep in re.findall(r"[,.]", whole)) or len(set(re.findall(r"[,.]", whole))) > 1:
			return None
		if not (1 <= len(groups[0]) <= 3 and all(len(g) in (2, 3) for g in groups[1:-1]) and len(groups[-1]) == 3):
			return None
	return float("".join(groups) + ("." + fraction if fraction else ""))


MONTHS = {
	name: index
	for index, names in enumerate(
		[
			("jan", "january"),
			("feb", "february"),
			("mar", "march"),
			("apr", "april"),
			("may",),
			("jun", "june"),
			("jul", "july"),
			("aug", "august"),
			("sep", "sept", "september"),
			("oct", "october"),
			("nov", "november"),
			("dec", "december"),
		],
		start=1,
	)
	for name in names
}


def parse_date(value, order: str = "DMY") -> dict | None:
	"""A date cell as `{printed, year, month, day}`, or None if it is not one.

	Excel dates arrive as `datetime` and are taken as they are. Text is split
	into its three parts and put in `order`, since `03/04/2026` is March or April
	depending on the bank. A month written as a word is recognised wherever it
	stands. Nothing is checked against a calendar here: the browser does that
	when it converts, and a Bikram Sambat month can have 32 days.
	"""
	if value is None:
		return None
	if isinstance(value, datetime.datetime | datetime.date):
		return {"printed": _cell_text(value), "year": value.year, "month": value.month, "day": value.day}
	printed = str(value).translate(DEVANAGARI).strip()
	parts = re.findall(r"[A-Za-z]+|\d+", printed)
	if len(parts) < 3:
		return None
	parts = parts[:3]

	word = next((i for i, part in enumerate(parts) if part.isalpha()), None)
	if word is not None:
		month = MONTHS.get(parts[word].lower())
		numbers = [int(part) for i, part in enumerate(parts) if i != word and part.isdigit()]
		if not month or len(numbers) != 2:
			return None
		# "05 Mar 2026", "Mar 5, 2026", "2026 Mar 05": the year is the number
		# too large to be a day, and otherwise the last one ("05-Mar-26").
		first, second = numbers
		year, day = (first, second) if first > 31 else (second, first)
	else:
		if not all(part.isdigit() for part in parts):
			return None
		numbers = [int(part) for part in parts]
		if len(parts[0]) == 4:
			order = "YMD"
		positions = {letter: index for index, letter in enumerate(order)}
		year, month, day = (numbers[positions[letter]] for letter in "YMD")

	if year < 100:
		year += 2000
	if not (1 <= month <= 12 and 1 <= day <= 32 and 1900 <= year <= 2200):
		return None
	return {"printed": printed, "year": year, "month": month, "day": day}
